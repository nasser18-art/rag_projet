"""Validation, extraction et découpage des documents importés.

Formats réellement pris en charge — et testés : PDF (pypdf), DOCX
(python-docx), CSV (bibliothèque standard), XLSX (openpyxl). Toute autre
extension est refusée explicitement (415) : l'interface ne doit jamais annoncer
un format que ce module ne sait pas lire.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

ALLOWED_EXTENSIONS = ("pdf", "docx", "csv", "xlsx")
TEXT_EXTENSIONS = ("csv",)
OOXML_EXTENSIONS = ("docx", "xlsx")
SIGNATURES = {
    "pdf": (b"%PDF",),
    "docx": (b"PK\x03\x04",),
    "xlsx": (b"PK\x03\x04",),
}


# --------------------------------------------------------------------------- #
# Erreurs typées : le statut HTTP est porté par l'exception
# --------------------------------------------------------------------------- #

class DocumentError(Exception):
    status_code = 400

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


class UnsupportedFormat(DocumentError):
    status_code = 415


class FileTooLarge(DocumentError):
    status_code = 413


class EmptyDocument(DocumentError):
    status_code = 422


class ExtractionFailed(DocumentError):
    status_code = 422


class DuplicateDocument(DocumentError):
    status_code = 409


# --------------------------------------------------------------------------- #
# Modèles
# --------------------------------------------------------------------------- #

@dataclass(frozen=True)
class Page:
    """Bloc de texte extrait, avec sa provenance lisible dans l'interface."""
    label: str
    text: str


@dataclass(frozen=True)
class Chunk:
    ordinal: int
    label: str
    text: str


# --------------------------------------------------------------------------- #
# Validation
# --------------------------------------------------------------------------- #

def extension_of(filename: str) -> str:
    """Extension normalisée, sans point ni casse. Chaîne vide si absente."""
    return Path(filename or "").suffix.lower().lstrip(".")


def sanitize_filename(filename: str) -> str:
    """Conserve le nom lisible d'origine, sans chemin ni caractère de contrôle."""
    name = (filename or "").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    name = name.lstrip(".")  # pas de fichier caché
    if not name:
        raise UnsupportedFormat("Le fichier doit avoir un nom exploitable.")
    if len(name) > 160:
        stem, dot, suffix = name.rpartition(".")
        name = (stem[:150] + (dot + suffix if dot else "")) if dot else name[:160]
    return name


def human_size(size: int) -> str:
    if size >= 1024 * 1024:
        return f"{size / (1024 * 1024):.0f} Mo"
    if size >= 1024:
        return f"{size / 1024:.0f} Ko"
    return f"{size} octets"


def validate_upload(filename: str, content: bytes, *, max_bytes: int) -> str:
    """Vérifie nom, taille, signature. Retourne l'extension retenue."""
    ext = extension_of(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedFormat(
            "Format non pris en charge. Formats acceptés : " + ", ".join(ALLOWED_EXTENSIONS).upper() + "."
        )
    if not content:
        raise EmptyDocument("Le fichier est vide.")
    if len(content) > max_bytes:
        raise FileTooLarge(f"Fichier trop volumineux : maximum {human_size(max_bytes)}.")
    for signature in SIGNATURES.get(ext, ()):
        if not content.startswith(signature):
            raise UnsupportedFormat(
                f"Le contenu ne correspond pas à un fichier {ext.upper()} valide."
            )
    return ext


# --------------------------------------------------------------------------- #
# Extraction
# --------------------------------------------------------------------------- #

def clean_text(text: str) -> str:
    """Compacte les espaces et sauts de ligne, sans perdre les paragraphes."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_pdf(content: bytes) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages: list[Page] = []
    for index, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        if text:
            pages.append(Page(label=f"Page {index}", text=text))
    if not pages:
        raise EmptyDocument("Ce PDF ne contient aucun texte extractible (document scanné ?).")
    return pages


def _read_docx(content: bytes) -> list[Page]:
    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(content))
    blocks: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    text = clean_text("\n".join(block for block in blocks if block.strip()))
    if not text:
        raise EmptyDocument("Ce document Word ne contient aucun texte.")
    return [Page(label="Document", text=text)]


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionFailed("Encodage du fichier texte non reconnu.")


def _read_csv(content: bytes) -> list[Page]:
    text = _decode_text(content)
    if "\x00" in text:
        raise UnsupportedFormat("Ce fichier ne semble pas être un CSV texte.")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    lines: list[str] = []
    for row in reader:
        cells = [cell.strip() for cell in row if cell and cell.strip()]
        if cells:
            lines.append(" | ".join(cells))
    body = clean_text("\n".join(lines))
    if not body:
        raise EmptyDocument("Ce CSV est vide.")
    return [Page(label="Tableau", text=body)]


def _read_xlsx(content: bytes) -> list[Page]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    pages: list[Page] = []
    for sheet in workbook.worksheets:
        lines: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if cells:
                lines.append(" | ".join(cells))
        text = clean_text("\n".join(lines))
        if text:
            pages.append(Page(label=f"Feuille {sheet.title}", text=text))
    workbook.close()
    if not pages:
        raise EmptyDocument("Ce classeur ne contient aucune cellule exploitable.")
    return pages


READERS = {
    "pdf": _read_pdf,
    "docx": _read_docx,
    "csv": _read_csv,
    "xlsx": _read_xlsx,
}


def extract_pages(content: bytes, ext: str) -> list[Page]:
    """Extrait les blocs de texte d'un fichier déjà validé."""
    reader = READERS.get(ext)
    if reader is None:
        raise UnsupportedFormat(f"Aucun extracteur pour l'extension « {ext} ».")
    try:
        return reader(content)
    except DocumentError:
        raise
    except Exception as exc:  # pypdf / openpyxl / docx lèvent des erreurs très variées
        raise ExtractionFailed(
            f"Lecture impossible du fichier {ext.upper()} : {type(exc).__name__}."
        ) from exc


# --------------------------------------------------------------------------- #
# Découpage
# --------------------------------------------------------------------------- #

def _split(text: str, chunk_size: int, overlap: int) -> list[str]:
    """Découpage récursif ; repli déterministe si langchain est absent."""
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter
    except ImportError:
        step = max(1, chunk_size - overlap)
        return [text[i:i + chunk_size] for i in range(0, len(text), step) if text[i:i + chunk_size].strip()]

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=overlap,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return splitter.split_text(text)


def chunk_pages(pages: Iterable[Page], *, chunk_size: int, chunk_overlap: int) -> list[Chunk]:
    """Découpe page par page pour conserver une provenance exacte."""
    chunks: list[Chunk] = []
    ordinal = 0
    for page in pages:
        for piece in _split(page.text, chunk_size, chunk_overlap):
            piece = piece.strip()
            if not piece:
                continue
            chunks.append(Chunk(ordinal=ordinal, label=page.label, text=piece))
            ordinal += 1
    if not chunks:
        raise EmptyDocument("Aucun contenu exploitable après découpage.")
    return chunks


def sha256_of(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()
