# ============================================================
# 03_test_extraction_docx.py
# Objectif : extraire le texte d'UN SEUL fichier Word (.docx)
# ============================================================

from pathlib import Path
from docx import Document

# --- 1. Remplacez par le chemin d'un de vos 2 fichiers Word ---
# Exemple : dans le dossier "informatique" ou "rh"
CHEMIN_DOCX = Path(__file__).resolve().parent / "documents" / "informatique" / "charte_informatique.docx"

# --- 2. Vérification ---
if not CHEMIN_DOCX.exists():
    print(f"❌ Fichier introuvable : {CHEMIN_DOCX}")
    print("👉 Vérifiez le nom exact du fichier dans le dossier.")
    exit()

print(f"📄 Fichier : {CHEMIN_DOCX.name}")

# --- 3. Ouverture du Word ---
doc = Document(str(CHEMIN_DOCX))

# --- 4. Extraction des paragraphes ---
texte_complet = ""
for para in doc.paragraphs:
    texte_complet += para.text + "\n"

print(f"📏 Nombre de paragraphes : {len(doc.paragraphs)}")
print(f"📏 Taille totale du texte extrait : {len(texte_complet)} caractères")
print()

# --- 5. Aperçu ---
print("=" * 60)
print("🔍 APERÇU DU DÉBUT (500 premiers caractères)")
print("=" * 60)
print(texte_complet[:500])
print()
print("=" * 60)
print("🔍 APERÇU DE LA FIN (500 derniers caractères)")
print("=" * 60)
print(texte_complet[-500:])
