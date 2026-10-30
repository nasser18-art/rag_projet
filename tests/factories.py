"""Fabriques de fichiers de test et encodeur déterministe.

Aucun fichier binaire n'est versionné : les documents de test sont générés à la
volée (PDF écrit à la main, DOCX via python-docx, XLSX via openpyxl, CSV en
texte), et l'encodeur d'embeddings est remplacé par une fonction déterministe —
le vrai modèle n'est jamais nécessaire pour vérifier la logique.
"""

from __future__ import annotations

import io
import re

import numpy as np


def make_pdf(pages: list[str]) -> bytes:
    """Construit un PDF 1.4 minimal, lisible par pypdf."""
    font_id = 3
    page_ids = [4 + 2 * index for index in range(len(pages))]
    content_ids = [5 + 2 * index for index in range(len(pages))]

    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: (
            f"<< /Type /Pages /Kids [{' '.join(f'{pid} 0 R' for pid in page_ids)}] "
            f"/Count {len(pages)} >>"
        ).encode(),
        font_id: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }

    for text, page_id, content_id in zip(pages, page_ids, content_ids):
        escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        stream = f"BT /F1 11 Tf 40 760 Td 12 TL ({escaped}) Tj ET".encode("latin-1", "replace")
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        objects[content_id] = (
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )

    output = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for object_id in sorted(objects):
        offsets[object_id] = len(output)
        output += f"{object_id} 0 obj\n".encode() + objects[object_id] + b"\nendobj\n"

    xref_position = len(output)
    largest = max(objects)
    output += f"xref\n0 {largest + 1}\n".encode() + b"0000000000 65535 f \n"
    for object_id in range(1, largest + 1):
        output += f"{offsets[object_id]:010d} 00000 n \n".encode()
    output += (
        f"trailer\n<< /Size {largest + 1} /Root 1 0 R >>\nstartxref\n{xref_position}\n%%EOF\n".encode()
    )
    return bytes(output)


def make_docx(paragraphs: list[str]) -> bytes:
    from docx import Document as DocxDocument

    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_csv(rows: list[list[str]]) -> bytes:
    return ("\n".join(";".join(row) for row in rows) + "\n").encode("utf-8")


def make_xlsx(rows: list[list[object]], sheet_title: str = "Fournisseurs") -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


TOKEN = re.compile(r"[\wàâäéèêëîïôöùûüçÀÂÄÉÈÊËÎÏÔÖÙÛÜÇ]+")


class FakeEncoder:
    """Sac de mots projeté aléatoirement, normalisé : similarité cosinus réelle."""

    def __init__(self, dimensions: int = 256) -> None:
        self.dimensions = dimensions
        self.calls = 0

    def _vector(self, text: str) -> np.ndarray:
        vector = np.zeros(self.dimensions, dtype=np.float32)
        for token in TOKEN.findall(text.lower()):
            vector[abs(hash(token)) % self.dimensions] += 1.0
        norm = np.linalg.norm(vector)
        return vector / norm if norm else vector

    def __call__(self, texts: list[str]) -> np.ndarray:
        self.calls += len(texts)
        return np.vstack([self._vector(text) for text in texts])
